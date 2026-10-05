"""
Export utilities for saving scraped Avito items to Excel, CSV, JSON, and DataFrames.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from avito_sdk.models import Item


def to_json(
    items: List[Item],
    filepath: Union[str, Path],
    indent: int = 2,
    include_raw: bool = False,
) -> None:
    """Save items to a formatted JSON file."""
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = [item.to_dict(include_raw=include_raw) for item in items]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=indent)


def to_jsonl(
    items: List[Item],
    filepath: Union[str, Path],
    include_raw: bool = False,
) -> None:
    """Save items to a JSON Lines (JSONL) file, ideal for streaming datasets."""
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        for item in items:
            f.write(item.to_json(include_raw=include_raw) + "\n")


def to_csv(
    items: List[Item],
    filepath: Union[str, Path],
    delimiter: str = ";",
) -> None:
    """Save items to a standard CSV file."""
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = [
        "id",
        "title",
        "price",
        "old_price",
        "price_drop",
        "seller_name",
        "seller_id",
        "url",
        "location_name",
        "address",
        "category_name",
        "total_views",
        "today_views",
        "is_promotion",
        "published_at",
        "params",
        "description",
    ]

    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=delimiter)
        writer.writeheader()
        for item in items:
            d = item.to_dict()
            # Format params dictionary as readable text
            params_str = "\n".join(f"{k}: {v}" for k, v in (item.params or {}).items())
            writer.writerow({
                "id": d["id"],
                "title": d["title"],
                "price": d["price"],
                "old_price": d.get("old_price"),
                "price_drop": d.get("price_drop"),
                "seller_name": d.get("seller_name"),
                "seller_id": d.get("seller_id"),
                "url": d["url"],
                "location_name": d.get("location_name"),
                "address": d.get("address"),
                "category_name": d.get("category_name"),
                "total_views": d.get("total_views"),
                "today_views": d.get("today_views"),
                "is_promotion": d.get("is_promotion"),
                "published_at": d.get("published_at"),
                "params": params_str,
                "description": (d.get("description") or "").replace("\n", " "),
            })


def to_excel(
    items: List[Item],
    filepath: Union[str, Path],
    sheet_name: str = "Avito Listings",
) -> None:
    """
    Save items to a styled Excel (.xlsx) file using openpyxl.
    Gracefully falls back to CSV if openpyxl is not installed.
    """
    path = Path(filepath)
    path.parent.mkdir(parents=True, exist_ok=True)

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        # Fallback to CSV if openpyxl is missing
        csv_path = path.with_suffix(".csv")
        to_csv(items, csv_path)
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        "ID",
        "Название",
        "Цена (₽)",
        "Старая цена (₽)",
        "Снижение (₽)",
        "Продавец",
        "ID продавца",
        "Ссылка",
        "Город / Регион",
        "Адрес",
        "Категория",
        "Просмотров всего",
        "Просмотров сегодня",
        "Продвижение",
        "Дата публикации",
        "Характеристики",
        "Описание",
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for item in items:
        params_str = "\n".join(f"{k}: {v}" for k, v in (item.params or {}).items())
        pub_str = item.published_at.strftime("%Y-%m-%d %H:%M") if item.published_at else ""

        row = [
            item.id,
            item.title,
            item.price,
            item.old_price,
            item.price_drop,
            item.seller_name or "",
            item.seller_id or "",
            item.url,
            item.location_name or "",
            item.address or "",
            item.category_name or "",
            item.total_views,
            item.today_views,
            "Да" if item.is_promotion else "Нет",
            pub_str,
            params_str,
            item.description or "",
        ]
        ws.append(row)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(path)


def to_dataframe(items: List[Item]) -> Any:
    """Convert items to a pandas or polars DataFrame if installed."""
    data = [item.to_dict() for item in items]
    try:
        import pandas as pd
        return pd.DataFrame(data)
    except ImportError:
        pass

    try:
        import polars as pl
        return pl.DataFrame(data)
    except ImportError:
        pass

    return data
